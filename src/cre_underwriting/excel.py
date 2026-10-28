"""Read deal assumptions from the Excel model's input block."""

from pathlib import Path

from openpyxl import load_workbook

from .assumptions import DealAssumptions

# Row label in column A -> DealAssumptions field. Matching on labels (not cell
# addresses) means inserting a row in the spreadsheet won't silently break this.
LABEL_TO_FIELD = {
    "Purchase price": "purchase_price",
    "Closing costs (% of price)": "closing_cost_pct",
    "Building size (SF)": "building_sf",
    "Starting rent ($/SF/yr)": "rent_per_sf",
    "Annual rent growth": "rent_growth",
    "Vacancy & credit loss": "vacancy_pct",
    "Operating expenses ($/SF/yr)": "opex_per_sf",
    "Expense growth": "expense_growth",
    "Capex reserve ($/SF/yr)": "capex_per_sf",
    "Loan-to-value (LTV)": "ltv",
    "Interest rate": "interest_rate",
    "Exit cap rate": "exit_cap_rate",
    "Selling costs (% of sale price)": "selling_cost_pct",
}


def load_assumptions(path: str | Path, sheet: str = "Model") -> DealAssumptions:
    ws = load_workbook(path, data_only=True)[sheet]
    values: dict[str, float] = {}
    for label, value in ws.iter_rows(min_col=1, max_col=2, values_only=True):
        field = LABEL_TO_FIELD.get(label)
        # The sensitivity table reuses "Exit cap rate" as a header; keep the first (input) one.
        if field and field not in values:
            values[field] = value

    missing = set(LABEL_TO_FIELD.values()) - values.keys()
    if missing:
        raise ValueError(f"Inputs not found in {path}: {sorted(missing)}")
    return DealAssumptions(**values)
